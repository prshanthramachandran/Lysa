"""
Root Growth analysis — specialised workflow for seedling root tracing on
scanned agar plates.

Workflow:
    1. User draws a rectangle per plate (ROI-style).
    2. User draws a rectangle per genotype region inside each plate.
    3. For each root: user clicks the shoot (start), then each timepoint
       marker in order (T24H, T48H, ...). Backend thresholds the plate crop,
       skeletonises, and returns a shortest-path polyline from shoot to the
       last timepoint click, passing through each click.
    4. Frontend lets the user edit the polyline; on save, backend computes
       per-segment lengths (shoot→T24H, T24H→T48H, ...) and persists the
       root entry as a JSON sidecar.

All length calculations are returned in both pixels and the image's physical
unit (uses metadata.pixel_size_x — typically µm from TIFF DPI). An Excel
export wraps up the results per plate/genotype/root.
"""

from __future__ import annotations

import json
import math
import uuid
from pathlib import Path
from typing import List, Optional, Tuple

import numpy as np
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from .. import store

router = APIRouter(prefix="/api/root-growth", tags=["root-growth"])

# Sidecar persistence directory
_SIDE_DIR = Path(__file__).resolve().parent.parent.parent / "uploads" / "_root_growth"
_SIDE_DIR.mkdir(parents=True, exist_ok=True)


# ---------------------------------------------------------------------------
# Models
# ---------------------------------------------------------------------------

class Rect(BaseModel):
    x: float
    y: float
    w: float
    h: float


class Point(BaseModel):
    x: float
    y: float


class TracePoint(BaseModel):
    """A click with its role label (e.g. SHOOT, T24H, T48H...)."""
    x: float
    y: float
    label: str


class AutoTraceParams(BaseModel):
    image_id: str
    plate_rect: Rect
    points: List[TracePoint]          # [SHOOT, T24H, T48H, ...] in order
    threshold_method: str = "otsu"    # "otsu" | "manual"
    manual_threshold: int = 128
    invert: bool = True               # dark roots on light agar → invert
    smooth_sigma: float = 1.0
    min_object_size: int = 50
    remove_rim: bool = True           # drop plate rim from the mask; preserves
                                      # components under user click points


class TraceFromPointParams(BaseModel):
    image_id: str
    plate_rect: Rect
    x: float                          # one click near the root's top (image coords)
    y: float
    threshold_method: str = "otsu"
    manual_threshold: int = 128
    invert: bool = True
    smooth_sigma: float = 1.0
    min_object_size: int = 50
    snap_radius: int = 60             # px to search for the skeleton near the click


class RootSegment(BaseModel):
    label: str              # "shoot→T24H"
    from_label: str
    to_label: str
    length_px: float
    length_physical: Optional[float] = None
    physical_unit: Optional[str] = None


class RootEntry(BaseModel):
    root_id: str
    image_id: str
    plate_index: int = 0
    plate_rect: Rect
    genotype: str
    polyline: List[Point]        # ordered (x,y) in image coordinates
    timepoints: List[TracePoint] # SHOOT + T24H + T48H ...
    segments: List[RootSegment] = Field(default_factory=list)
    total_length_px: float = 0.0
    total_length_physical: Optional[float] = None
    physical_unit: Optional[str] = None
    tortuosity: Optional[float] = None
    angle_deg: Optional[float] = None
    deviation_deg: Optional[float] = None


class SaveSessionParams(BaseModel):
    image_id: str
    plates: List[Rect]
    genotype_names: List[str]
    genotype_boxes: List[Rect]          # parallel to genotype_names
    genotype_plate_idx: List[int] = Field(default_factory=list)
    timepoint_labels: List[str]         # e.g. ["SHOOT","T24H","T48H"]
    roots: List[RootEntry]


class ComputeSegmentsParams(BaseModel):
    image_id: str
    polyline: List[Point]
    timepoints: List[TracePoint]        # in traversal order


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _pixel_size(image_id: str) -> Tuple[Optional[float], Optional[str]]:
    entry = store.get(image_id)
    meta = entry.get("metadata", {})
    return meta.get("pixel_size_x"), meta.get("pixel_size_unit")


# ---------------------------------------------------------------------------
# Calibration (physical scale: µm per pixel)
# ---------------------------------------------------------------------------

_UNIT_TO_UM = {"mm": 1000.0, "µm": 1.0, "um": 1.0, "cm": 10000.0, "nm": 0.001}


@router.get("/calibration/{image_id}")
def get_calibration(image_id: str):
    """Report the active physical scale for an image and where it came from.

    source is one of:
      "manual"  – set via /calibrate (ruler), takes precedence,
      "dpi"     – derived from the image's TIFF resolution tags at ingest,
      "none"    – uncalibrated; lengths will be in PIXELS only.
    """
    if not store.contains(image_id):
        raise HTTPException(404, "Unknown image_id")
    meta = store.get(image_id).get("metadata", {})
    px = meta.get("pixel_size_x")
    unit = meta.get("pixel_size_unit")
    source = meta.get("pixel_size_source") or ("dpi" if px else "none")
    return {
        "pixel_size_x": px,
        "pixel_size_unit": unit,
        "calibrated": px is not None,
        "source": source,
    }


class CalibrateParams(BaseModel):
    image_id: str
    x1: float
    y1: float
    x2: float
    y2: float
    known_distance: float            # real-world distance between the 2 points
    unit: str = "mm"                 # unit of known_distance (mm/µm/cm/nm)


@router.post("/calibrate")
def calibrate(params: CalibrateParams):
    """Set an image's scale from two clicked points a known distance apart.

    Computes µm/pixel = (known_distance in µm) / (pixel distance between the
    two points), stores it on the image metadata (as pixel_size_x/y in µm,
    source="manual"), and returns the new scale. Used with the in-frame ruler
    when DPI is missing or you want to override it.
    """
    if not store.contains(params.image_id):
        raise HTTPException(404, "Unknown image_id")
    dist_px = math.hypot(params.x2 - params.x1, params.y2 - params.y1)
    if dist_px < 1.0:
        raise HTTPException(400, "The two points are too close together.")
    if params.known_distance <= 0:
        raise HTTPException(400, "known_distance must be positive.")
    unit = params.unit if params.unit in _UNIT_TO_UM else "mm"
    known_um = params.known_distance * _UNIT_TO_UM[unit]
    um_per_px = round(known_um / dist_px, 6)

    meta = store.get(params.image_id).get("metadata", {})
    meta["pixel_size_x"] = um_per_px
    meta["pixel_size_y"] = um_per_px
    meta["pixel_size_unit"] = "µm"
    meta["pixel_size_source"] = "manual"
    return {
        "pixel_size_x": um_per_px,
        "pixel_size_unit": "µm",
        "calibrated": True,
        "source": "manual",
        "pixel_distance": round(dist_px, 2),
    }


def _polyline_length_px(pts: List[Point]) -> float:
    if len(pts) < 2:
        return 0.0
    total = 0.0
    for a, b in zip(pts[:-1], pts[1:]):
        total += math.hypot(b.x - a.x, b.y - a.y)
    return total


def _root_metrics(poly: List[Point]) -> dict:
    """Shape metrics for a traced root polyline.

    Returns:
      tortuosity      path length / straight-line (chord) distance, >= 1.0.
                      1.0 = perfectly straight; higher = more wandering.
                      None if the chord is ~0 (degenerate).
      chord_px        straight-line distance from first to last point (px).
      angle_deg       growth direction of the chord, in DEGREES, in image
                      coordinates where +y points DOWN (as in pixel space).
                      0 = pointing right (+x), 90 = straight down (+y, the
                      usual gravitropic "down the plate" direction), 180 =
                      left, -90 = up. Measured tip-relative: from the first
                      point (shoot) to the last (root tip).
      deviation_deg   absolute deviation from straight-down (90°), i.e. how
                      far the root strayed from the gravity vector. 0 = grew
                      straight down; 90 = grew horizontally.
    """
    if len(poly) < 2:
        return {"tortuosity": None, "chord_px": 0.0,
                "angle_deg": None, "deviation_deg": None}
    dx = poly[-1].x - poly[0].x
    dy = poly[-1].y - poly[0].y
    chord = math.hypot(dx, dy)
    path_len = _polyline_length_px(poly)
    tort = round(path_len / chord, 4) if chord > 1e-6 else None
    if chord <= 1e-6:
        return {"tortuosity": tort, "chord_px": round(chord, 3),
                "angle_deg": None, "deviation_deg": None}
    angle = math.degrees(math.atan2(dy, dx))   # +y down → +90 = down
    deviation = abs(angle - 90.0)
    if deviation > 180.0:
        deviation = 360.0 - deviation
    return {
        "tortuosity": tort,
        "chord_px": round(chord, 3),
        "angle_deg": round(angle, 2),
        "deviation_deg": round(deviation, 2),
    }


def _project_point_onto_polyline(
    px: float, py: float, poly: List[Point]
) -> Tuple[int, float, float]:
    """
    Project (px, py) onto the polyline. Returns (segment_index, t, cum_len)
    where `t` ∈ [0,1] is the fraction along the segment and `cum_len` is the
    distance along the polyline to the projection.
    """
    best = (0, 0.0, 0.0, float("inf"))
    cum = 0.0
    for i in range(len(poly) - 1):
        ax, ay = poly[i].x, poly[i].y
        bx, by = poly[i + 1].x, poly[i + 1].y
        dx, dy = bx - ax, by - ay
        seg_len_sq = dx * dx + dy * dy
        if seg_len_sq == 0:
            t = 0.0
            qx, qy = ax, ay
        else:
            t = ((px - ax) * dx + (py - ay) * dy) / seg_len_sq
            t = max(0.0, min(1.0, t))
            qx = ax + t * dx
            qy = ay + t * dy
        d = (px - qx) ** 2 + (py - qy) ** 2
        seg_len = math.sqrt(seg_len_sq)
        if d < best[3]:
            best = (i, t, cum + t * seg_len, d)
        cum += seg_len
    return best[0], best[1], best[2]


def _split_polyline_at_distances(
    poly: List[Point], distances: List[float]
) -> List[List[Point]]:
    """
    Split a polyline into N+1 sub-polylines given N cumulative distances
    (assumed sorted ascending). Each sub-polyline begins at the previous
    split point.
    """
    segs: List[List[Point]] = []
    current: List[Point] = [poly[0]]
    cum = 0.0
    dist_iter = iter(distances)
    next_d = next(dist_iter, None)

    for i in range(len(poly) - 1):
        a, b = poly[i], poly[i + 1]
        seg_len = math.hypot(b.x - a.x, b.y - a.y)
        seg_start = cum
        seg_end = cum + seg_len
        # Inject any split points that fall within this segment
        while next_d is not None and next_d <= seg_end + 1e-9:
            if seg_len > 0:
                t = (next_d - seg_start) / seg_len
                t = max(0.0, min(1.0, t))
            else:
                t = 0.0
            qx = a.x + t * (b.x - a.x)
            qy = a.y + t * (b.y - a.y)
            current.append(Point(x=qx, y=qy))
            segs.append(current)
            current = [Point(x=qx, y=qy)]
            next_d = next(dist_iter, None)
        current.append(b)
        cum = seg_end

    segs.append(current)
    return segs


def _build_segments(
    polyline: List[Point], timepoints: List[TracePoint], image_id: str
) -> Tuple[List[RootSegment], float, float, Optional[str]]:
    """
    Project each timepoint onto the polyline, split it, return segments with
    per-segment length and total length in pixels + physical units.
    """
    if len(polyline) < 2 or len(timepoints) < 2:
        return [], 0.0, 0.0, None

    px_size, unit = _pixel_size(image_id)

    # Project every timepoint → cumulative distance along the polyline
    projections = []
    for tp in timepoints:
        _, _, cum = _project_point_onto_polyline(tp.x, tp.y, polyline)
        projections.append((cum, tp.label))

    # Sort by cum distance to get traversal order
    projections.sort(key=lambda p: p[0])

    total_len_px = _polyline_length_px(polyline)
    segments: List[RootSegment] = []
    for i in range(len(projections) - 1):
        d0, l0 = projections[i]
        d1, l1 = projections[i + 1]
        length_px = max(0.0, d1 - d0)
        length_phys = round(length_px * px_size, 4) if px_size else None
        segments.append(
            RootSegment(
                label=f"{l0}→{l1}",
                from_label=l0,
                to_label=l1,
                length_px=round(length_px, 3),
                length_physical=length_phys,
                physical_unit=unit,
            )
        )

    total_px = sum(s.length_px for s in segments)
    total_phys = round(total_px * px_size, 4) if px_size else None
    return segments, total_px, total_phys, unit


# ---------------------------------------------------------------------------
# Auto plate detection — find rectangular plates without manual drawing
# ---------------------------------------------------------------------------

class DetectPlatesParams(BaseModel):
    image_id: str
    expected: Optional[int] = None     # hint: expected plate count (optional)
    min_area_frac: float = 0.03        # ignore blobs < this fraction of image
    # Plates are usually the BRIGHT agar regions on the scanner bed. Set
    # dark_plates=True for the rarer case of dark plates on a light bed.
    dark_plates: bool = False
    sep_strength: int = 9              # erosion kernel — higher = separate
                                       # plates that touch/bridge more aggressively


@router.post("/detect-plates")
def detect_plates(params: DetectPlatesParams):
    """Auto-detect square-ish plate regions on a scan.

    Real agar-plate scans are typically BRIGHT plates on a bright scanner bed,
    separated only by thin dark rims — so a plain bright/dark threshold merges
    everything into one blob. Instead we:

      1. downscale for speed,
      2. threshold to the bright (agar) class — or dark, if dark_plates,
      3. **erode** hard so thin bridges through rims / background between
         plates are severed, isolating each plate interior as its own blob,
      4. keep large, square-ish components (rejects the ruler / colour strip /
         slivers), and
      5. dilate each bounding box back out to undo the erosion shrinkage.

    Boxes are returned in FULL-RESOLUTION coords, sorted row-major. They drop
    straight into the editable per-plate rectangles, so this augments — never
    replaces — the manual workflow. Tunable via dark_plates / sep_strength /
    min_area_frac if a particular scan needs it.
    """
    import numpy as np
    from skimage.filters import threshold_otsu
    from skimage.morphology import binary_opening, remove_small_objects, disk
    from scipy import ndimage as ndi

    if not store.contains(params.image_id):
        raise HTTPException(404, "Unknown image_id")

    display = store.load_display_array(params.image_id)
    gray = _to_gray(display)
    H, W = gray.shape[:2]

    # Downscale (integer stride both ways so boxes map back exactly).
    step = max(1, int(max(H, W) / 1000.0))
    small = gray[::step, ::step]
    sh, sw = small.shape[:2]
    img_area = sh * sw
    min_px = int(params.min_area_frac * img_area)

    try:
        t = float(threshold_otsu(small))
    except Exception:
        t = 128.0

    # Class of interest: bright agar (default) or dark plates, split at Otsu.
    # (A margin below t was tried to keep dim agar, but it pulls a bright
    # scanner background into the bright class and merges everything — plain
    # Otsu separates plate-vs-bed cleanly on both real scans and synthetic.)
    region = small < t if params.dark_plates else small > t
    region = binary_opening(region, disk(2))

    # Sever bridges (rims, background necks, ruler stems) between plates.
    k = max(3, int(params.sep_strength))
    eroded = ndi.binary_erosion(region, structure=np.ones((k, k)))
    try:
        eroded = remove_small_objects(eroded, min_size=max(1, min_px))
    except TypeError:
        eroded = remove_small_objects(eroded, max(1, min_px))

    labels, n = ndi.label(eroded)
    if n == 0:
        return {"plates": [], "count": 0,
                "note": "No plate-like regions found. Try toggling "
                        "'plates darker than background', or draw manually."}

    pad = k // 2  # dilate boxes back by ~the erosion radius
    boxes = []
    for sl in ndi.find_objects(labels):
        if sl is None:
            continue
        ys, xs = sl
        bw, bh = xs.stop - xs.start, ys.stop - ys.start
        area = bw * bh
        if area < min_px:
            continue
        aspect = bw / bh if bh else 999
        if aspect < 0.4 or aspect > 2.5:        # plates are roughly square
            continue
        if area > 0.95 * img_area:              # whole-image background blob
            continue
        x0 = max(0, xs.start - pad)
        y0 = max(0, ys.start - pad)
        boxes.append((x0, y0,
                      min(sw - x0, bw + 2 * pad),
                      min(sh - y0, bh + 2 * pad)))

    # Suppress "merged" super-boxes: when erosion fails to fully sever two
    # plates, one big blob spanning both survives ALONGSIDE the two correctly
    # separated plate blobs. Drop any box that mostly covers a smaller box
    # (its children) — keep the tighter individual plates.
    def _contains(big, small, frac=0.75):
        bx, by, bw, bh = big
        sx, sy, sw, sh = small
        ix0, iy0 = max(bx, sx), max(by, sy)
        ix1, iy1 = min(bx + bw, sx + sw), min(by + bh, sy + sh)
        inter = max(0, ix1 - ix0) * max(0, iy1 - iy0)
        small_area = sw * sh
        return small_area > 0 and inter / small_area >= frac

    boxes.sort(key=lambda b: b[2] * b[3])  # small → large
    kept = []
    for b in boxes:
        # If b engulfs any already-kept (smaller) box, b is a merged super-box.
        if any(_contains(b, s) for s in kept):
            continue
        kept.append(b)
    boxes = kept

    # Map back to full-res, sort row-major (bucket rows by ~half median height).
    boxes = [(x * step, y * step, w * step, h * step) for (x, y, w, h) in boxes]
    if boxes:
        med_h = sorted(b[3] for b in boxes)[len(boxes) // 2]
        row_tol = max(1, med_h // 2)
        boxes.sort(key=lambda b: (b[1] // row_tol, b[0]))

    plates = [{"x": x, "y": y, "w": w, "h": h} for (x, y, w, h) in boxes]
    note = None
    if params.expected and len(plates) != params.expected:
        note = (f"Found {len(plates)} plate(s), expected {params.expected}. "
                f"Try the 'plates darker than background' toggle, adjust "
                f"separation strength, or edit the rectangles manually.")
    return {"plates": plates, "count": len(plates), "note": note}


# ---------------------------------------------------------------------------
# Auto-trace — threshold → skeletonise → shortest path
# ---------------------------------------------------------------------------

def _crop_plate(arr: np.ndarray, rect: Rect) -> Tuple[np.ndarray, int, int]:
    h, w = arr.shape[:2]
    x0 = max(0, int(round(rect.x)))
    y0 = max(0, int(round(rect.y)))
    x1 = min(w, int(round(rect.x + rect.w)))
    y1 = min(h, int(round(rect.y + rect.h)))
    if x1 <= x0 or y1 <= y0:
        raise HTTPException(400, "Plate rectangle is empty / out of bounds")
    return arr[y0:y1, x0:x1].copy(), x0, y0


def _to_gray(arr: np.ndarray) -> np.ndarray:
    if arr.ndim == 2:
        return arr.astype(np.float32)
    return np.mean(arr[..., :3].astype(np.float32), axis=2)


def _auto_trace_shortest_path(
    mask: np.ndarray, start: Tuple[int, int], end: Tuple[int, int]
) -> List[Tuple[int, int]]:
    """
    BFS over the binary mask from start to end. Returns list of (y,x).
    If disconnected, returns empty list.
    """
    from collections import deque
    h, w = mask.shape
    sy, sx = start
    ey, ex = end
    if not (0 <= sy < h and 0 <= sx < w and 0 <= ey < h and 0 <= ex < w):
        return []
    # Snap start/end to nearest mask pixel if not already on the mask
    def _snap(y, x):
        if mask[y, x]:
            return y, x
        # Ring search
        for r in range(1, 40):
            y0 = max(0, y - r); y1 = min(h, y + r + 1)
            x0 = max(0, x - r); x1 = min(w, x + r + 1)
            sub = mask[y0:y1, x0:x1]
            if sub.any():
                ys, xs = np.where(sub)
                return ys[0] + y0, xs[0] + x0
        return y, x

    sy, sx = _snap(sy, sx)
    ey, ex = _snap(ey, ex)

    if not mask[sy, sx] or not mask[ey, ex]:
        return []

    prev = -np.ones((h, w), dtype=np.int64)
    prev[sy, sx] = sy * w + sx
    q = deque()
    q.append((sy, sx))
    found = False
    # 8-connected
    steps = [(-1, -1), (-1, 0), (-1, 1), (0, -1), (0, 1), (1, -1), (1, 0), (1, 1)]
    while q:
        y, x = q.popleft()
        if y == ey and x == ex:
            found = True
            break
        for dy, dx in steps:
            ny, nx = y + dy, x + dx
            if 0 <= ny < h and 0 <= nx < w and mask[ny, nx] and prev[ny, nx] == -1:
                prev[ny, nx] = y * w + x
                q.append((ny, nx))
    if not found:
        return []
    # Reconstruct
    path = []
    y, x = ey, ex
    while True:
        path.append((y, x))
        p = prev[y, x]
        py, px = divmod(int(p), w)
        if (py, px) == (y, x):
            break
        y, x = py, px
    path.reverse()
    return path


def _simplify_polyline(pts: List[Tuple[int, int]], step: int = 8) -> List[Tuple[int, int]]:
    """Subsample every `step`-th pixel; always keep endpoints."""
    if len(pts) <= 2:
        return pts
    out = pts[::step]
    if out[-1] != pts[-1]:
        out.append(pts[-1])
    return out


def _trace_down_skeleton(skel, start):
    """From `start` (y,x) on the skeleton, return the path to the geodesically
    farthest reachable skeleton pixel that lies at or below the start row.

    Used by single-click tracing: the user clicks the shoot, and the root is
    the longest downward run of skeleton from there. Returns list of (y,x).
    """
    from collections import deque
    h, w = skel.shape
    sy, sx = start
    prev = -np.ones((h, w), dtype=np.int64)
    dist = -np.ones((h, w), dtype=np.float64)
    prev[sy, sx] = sy * w + sx
    dist[sy, sx] = 0.0
    q = deque([(sy, sx)])
    steps = [(-1, -1), (-1, 0), (-1, 1), (0, -1), (0, 1), (1, -1), (1, 0), (1, 1)]
    far, fard = (sy, sx), 0.0
    while q:
        y, x = q.popleft()
        for dy, dx in steps:
            ny, nx = y + dy, x + dx
            if 0 <= ny < h and 0 <= nx < w and skel[ny, nx] and prev[ny, nx] == -1:
                prev[ny, nx] = y * w + x
                dist[ny, nx] = dist[y, x] + (1.4142 if dy and dx else 1.0)
                q.append((ny, nx))
                if ny >= sy and dist[ny, nx] > fard:
                    fard, far = dist[ny, nx], (ny, nx)
    path = []
    y, x = far
    while True:
        path.append((y, x))
        pp = prev[y, x]
        py, px = divmod(int(pp), w)
        if (py, px) == (y, x):
            break
        y, x = py, px
    path.reverse()
    return path


@router.post("/trace-from-point")
def trace_from_point(params: TraceFromPointParams):
    """Trace one root from a SINGLE click — no tip click needed.

    Snap the click to the nearest skeleton pixel, then walk the skeleton to its
    farthest downward point (the root tip). Returns an editable polyline plus
    length + shape metrics. Best-effort: where a root's skeleton is broken by a
    faint gap the trace may stop short — the returned polyline is editable, so
    the user drags to extend. Vertical gap-bridging reduces this.
    """
    from skimage.filters import threshold_otsu
    from skimage.morphology import (
        remove_small_objects, skeletonize, binary_closing, disk
    )
    from skimage.segmentation import clear_border
    from scipy import ndimage as ndi

    if not store.contains(params.image_id):
        raise HTTPException(404, "Unknown image_id")

    display = store.load_display_array(params.image_id)
    crop, ox, oy = _crop_plate(display, params.plate_rect)
    gray = _to_gray(crop)
    if params.smooth_sigma > 0:
        gray = ndi.gaussian_filter(gray, sigma=float(params.smooth_sigma))
    if params.threshold_method == "otsu":
        try:
            t = float(threshold_otsu(gray))
        except Exception:
            t = 128.0
    else:
        t = float(params.manual_threshold)
    mask = gray < t if params.invert else gray > t
    mask = binary_closing(mask, disk(2))
    mask = remove_small_objects(mask, min_size=int(max(1, params.min_object_size)))
    # rim removal (always on here — single-click tracing has no click to protect
    # other than the one we snap below, and the rim only causes mis-snaps)
    H, W = mask.shape
    mb = int(0.05 * min(H, W))
    if mb > 0:
        mask[:mb, :] = mask[-mb:, :] = False
        mask[:, :mb] = mask[:, -mb:] = False
    mask = clear_border(mask)
    mask = remove_small_objects(mask, min_size=int(max(1, params.min_object_size)))
    # bridge small vertical gaps so a fragmented root stays one path
    mask = ndi.binary_closing(mask, structure=np.ones((9, 3)))

    skel = skeletonize(mask)

    cy = int(round(params.y - oy))
    cx = int(round(params.x - ox))

    # Snap the click to the nearest skeleton pixel within snap_radius.
    def _snap(y, x, R):
        if 0 <= y < H and 0 <= x < W and skel[y, x]:
            return (y, x)
        for r in range(1, R):
            y0, y1 = max(0, y - r), min(H, y + r + 1)
            x0, x1 = max(0, x - r), min(W, x + r + 1)
            sub = skel[y0:y1, x0:x1]
            if sub.any():
                ys, xs = np.where(sub)
                return (ys[0] + y0, xs[0] + x0)
        return None

    snapped = _snap(cy, cx, int(max(5, params.snap_radius)))
    if snapped is None:
        raise HTTPException(422, "No root found near the click. Click closer to "
                                 "a root, or adjust threshold/invert.")

    path = _trace_down_skeleton(skel, snapped)
    if len(path) < 3:
        raise HTTPException(422, "Traced path too short — click higher on the root.")

    path = _simplify_polyline(path, step=6)
    polyline = [{"x": float(x + ox), "y": float(y + oy)} for (y, x) in path]

    length_px = 0.0
    for a, b in zip(polyline[:-1], polyline[1:]):
        length_px += math.hypot(b["x"] - a["x"], b["y"] - a["y"])
    px_size, unit = _pixel_size(params.image_id)
    metrics = _root_metrics([Point(x=p["x"], y=p["y"]) for p in polyline])
    return {
        "polyline": polyline,
        "length_px": round(length_px, 3),
        "length_physical": round(length_px * px_size, 4) if px_size else None,
        "physical_unit": unit,
        "n_points": len(polyline),
        "tortuosity": metrics["tortuosity"],
        "angle_deg": metrics["angle_deg"],
        "deviation_deg": metrics["deviation_deg"],
    }


@router.post("/auto-trace")
def auto_trace(params: AutoTraceParams):
    """
    Threshold the plate crop, skeletonise, and BFS a path through all
    click points in order. Returns the full polyline in image coordinates.
    """
    from skimage.filters import threshold_otsu
    from skimage.morphology import (
        remove_small_objects, skeletonize, binary_closing, disk
    )
    from skimage.segmentation import clear_border
    from scipy import ndimage as ndi

    if not store.contains(params.image_id):
        raise HTTPException(404, "Unknown image_id")
    if len(params.points) < 2:
        raise HTTPException(400, "Need at least 2 points (shoot + one timepoint)")

    display = store.load_display_array(params.image_id)
    crop, ox, oy = _crop_plate(display, params.plate_rect)
    gray = _to_gray(crop)

    if params.smooth_sigma > 0:
        gray = ndi.gaussian_filter(gray, sigma=float(params.smooth_sigma))

    if params.threshold_method == "otsu":
        try:
            t = float(threshold_otsu(gray))
        except Exception:
            t = 128.0
    else:
        t = float(params.manual_threshold)

    mask = gray < t if params.invert else gray > t

    # Clean up
    mask = binary_closing(mask, disk(2))
    mask = remove_small_objects(mask, min_size=int(max(1, params.min_object_size)))

    # --- Rim removal (on by default) -----------------------------------------
    # The plate rim hugs the crop edge and otherwise dominates the mask (it was
    # 47-75% of root pixels on real scans). Inset the border + clear_border to
    # drop it. CRUCIAL: preserve any connected component that contains a user
    # click (e.g. a shoot clicked near the top edge) so tracing never loses its
    # start/end anchors to rim cleanup.
    if params.remove_rim:
        H, W = mask.shape
        # Label the full mask first so we can rescue click-bearing components.
        lbl_full, _ = ndi.label(mask)
        protect = set()
        for pt in params.points:
            py = int(round(pt.y - oy)); px = int(round(pt.x - ox))
            if 0 <= py < H and 0 <= px < W and lbl_full[py, px]:
                protect.add(int(lbl_full[py, px]))
            else:
                # Click not on a mask pixel — protect the nearest component in a
                # small neighbourhood so a slightly-off shoot click still holds.
                y0, y1 = max(0, py - 15), min(H, py + 16)
                x0, x1 = max(0, px - 15), min(W, px + 16)
                sub = lbl_full[y0:y1, x0:x1]
                vals = sub[sub > 0]
                if vals.size:
                    protect.add(int(np.bincount(vals).argmax()))
        protected_mask = np.isin(lbl_full, list(protect)) if protect else np.zeros_like(mask)

        mb = int(0.05 * min(H, W))
        rimless = mask.copy()
        if mb > 0:
            rimless[:mb, :] = rimless[-mb:, :] = False
            rimless[:, :mb] = rimless[:, -mb:] = False
        rimless = clear_border(rimless)
        rimless = remove_small_objects(rimless, min_size=int(max(1, params.min_object_size)))
        # Add the protected (click-bearing) components back in full.
        mask = rimless | protected_mask

    # Skeletonise so BFS walks along the 1-pixel centre line
    skel = skeletonize(mask)

    # Piecewise BFS through all points in order
    full_path: List[Tuple[int, int]] = []
    for i in range(len(params.points) - 1):
        p0 = params.points[i]
        p1 = params.points[i + 1]
        seg = _auto_trace_shortest_path(
            skel,
            (int(round(p0.y - oy)), int(round(p0.x - ox))),
            (int(round(p1.y - oy)), int(round(p1.x - ox))),
        )
        if not seg:
            # Fallback: straight line
            n = 30
            y0, x0 = int(round(p0.y - oy)), int(round(p0.x - ox))
            y1, x1 = int(round(p1.y - oy)), int(round(p1.x - ox))
            seg = [
                (int(y0 + (y1 - y0) * k / n), int(x0 + (x1 - x0) * k / n))
                for k in range(n + 1)
            ]
        if full_path and seg and seg[0] == full_path[-1]:
            seg = seg[1:]
        full_path.extend(seg)

    full_path = _simplify_polyline(full_path, step=6)

    # Convert back to image coordinates
    polyline = [{"x": float(x + ox), "y": float(y + oy)} for (y, x) in full_path]

    length_px = 0.0
    for a, b in zip(polyline[:-1], polyline[1:]):
        length_px += math.hypot(b["x"] - a["x"], b["y"] - a["y"])

    px_size, unit = _pixel_size(params.image_id)
    length_phys = round(length_px * px_size, 4) if px_size else None

    poly_pts = [Point(x=p["x"], y=p["y"]) for p in polyline]
    metrics = _root_metrics(poly_pts)
    chord_phys = (round(metrics["chord_px"] * px_size, 4)
                  if px_size and metrics["chord_px"] else None)

    return {
        "polyline": polyline,
        "length_px": round(length_px, 3),
        "length_physical": length_phys,
        "physical_unit": unit,
        "n_points": len(polyline),
        "threshold_used": round(t, 2),
        "tortuosity": metrics["tortuosity"],
        "chord_px": metrics["chord_px"],
        "chord_physical": chord_phys,
        "angle_deg": metrics["angle_deg"],
        "deviation_deg": metrics["deviation_deg"],
    }


# ---------------------------------------------------------------------------
# Segment computation (no auto-trace — for manually drawn polylines)
# ---------------------------------------------------------------------------

@router.post("/compute-segments")
def compute_segments(params: ComputeSegmentsParams):
    if not store.contains(params.image_id):
        raise HTTPException(404, "Unknown image_id")
    segs, total_px, total_phys, unit = _build_segments(
        params.polyline, params.timepoints, params.image_id
    )
    px_size, _ = _pixel_size(params.image_id)
    metrics = _root_metrics(params.polyline)
    chord_phys = (round(metrics["chord_px"] * px_size, 4)
                  if px_size and metrics["chord_px"] else None)
    return {
        "segments": [s.model_dump() for s in segs],
        "total_length_px": round(total_px, 3),
        "total_length_physical": total_phys,
        "physical_unit": unit,
        "tortuosity": metrics["tortuosity"],
        "chord_px": metrics["chord_px"],
        "chord_physical": chord_phys,
        "angle_deg": metrics["angle_deg"],
        "deviation_deg": metrics["deviation_deg"],
    }


# ---------------------------------------------------------------------------
# Persistence (JSON sidecar per image)
# ---------------------------------------------------------------------------

def _sidecar_path(image_id: str) -> Path:
    return _SIDE_DIR / f"{image_id}.json"


@router.post("/save-session")
def save_session(params: SaveSessionParams):
    if not store.contains(params.image_id):
        raise HTTPException(404, "Unknown image_id")

    # Recompute segments + totals for every root before writing
    processed_roots = []
    for r in params.roots:
        segs, total_px, total_phys, unit = _build_segments(
            r.polyline, r.timepoints, params.image_id
        )
        r.segments = segs
        r.total_length_px = round(total_px, 3)
        r.total_length_physical = total_phys
        r.physical_unit = unit
        # Recompute shape metrics server-side so the saved values are
        # authoritative (don't trust whatever the client sent).
        m = _root_metrics(r.polyline)
        r.tortuosity = m["tortuosity"]
        r.angle_deg = m["angle_deg"]
        r.deviation_deg = m["deviation_deg"]
        if not r.root_id:
            r.root_id = uuid.uuid4().hex[:8]
        processed_roots.append(r)

    payload = {
        "image_id": params.image_id,
        "plates": [p.model_dump() for p in params.plates],
        "genotype_names": params.genotype_names,
        "genotype_boxes": [b.model_dump() for b in params.genotype_boxes],
        "genotype_plate_idx": params.genotype_plate_idx,
        "timepoint_labels": params.timepoint_labels,
        "roots": [r.model_dump() for r in processed_roots],
    }
    _sidecar_path(params.image_id).write_text(json.dumps(payload, indent=2))
    return {"saved": True, "n_roots": len(processed_roots)}


@router.get("/load-session/{image_id}")
def load_session(image_id: str):
    p = _sidecar_path(image_id)
    if not p.exists():
        return {"image_id": image_id, "exists": False}
    try:
        data = json.loads(p.read_text())
    except Exception as exc:
        raise HTTPException(500, f"Corrupted sidecar: {exc}")
    data["exists"] = True
    return data


@router.delete("/session/{image_id}")
def delete_session(image_id: str):
    p = _sidecar_path(image_id)
    if p.exists():
        p.unlink()
    return {"deleted": True}


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

@router.get("/export/{image_id}")
def export_session(image_id: str):
    import io
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(500, "openpyxl not installed")

    p = _sidecar_path(image_id)
    if not p.exists():
        raise HTTPException(404, "No Root Growth session saved for this image")
    data = json.loads(p.read_text())

    wb = Workbook()

    # --- Metadata sheet ---
    ws_meta = wb.active
    ws_meta.title = "Metadata"
    entry = store.get(image_id) if store.contains(image_id) else None
    meta = (entry or {}).get("metadata", {}) if entry else {}
    rows = [
        ("Image", meta.get("filename", image_id)),
        ("Image ID", image_id),
        ("Width (px)", meta.get("width")),
        ("Height (px)", meta.get("height")),
        ("Pixel size X", meta.get("pixel_size_x")),
        ("Pixel size unit", meta.get("pixel_size_unit")),
        ("Plates", len(data.get("plates", []))),
        ("Genotypes", ", ".join(data.get("genotype_names", []))),
        ("Timepoints", ", ".join(data.get("timepoint_labels", []))),
        ("Roots measured", len(data.get("roots", []))),
    ]
    bold = Font(bold=True)
    for r in rows:
        ws_meta.append(r)
        ws_meta.cell(row=ws_meta.max_row, column=1).font = bold
    ws_meta.column_dimensions["A"].width = 22
    ws_meta.column_dimensions["B"].width = 40

    # --- Measurements sheet (wide format) ---
    ws = wb.create_sheet("Root Measurements")
    tp_labels = data.get("timepoint_labels", [])
    segment_cols = [f"{tp_labels[i]}→{tp_labels[i+1]}" for i in range(len(tp_labels) - 1)]

    headers = ["Root ID", "Plate", "Genotype"]
    for col in segment_cols:
        headers.append(f"{col} (px)")
        headers.append(f"{col} (phys)")
    headers += ["Total (px)", "Total (phys)", "Unit",
                "Tortuosity", "Angle (deg)", "Deviation from down (deg)"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = bold
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
        cell.alignment = Alignment(horizontal="center")

    for root in data.get("roots", []):
        row = [
            root.get("root_id", ""),
            root.get("plate_index", 0) + 1,
            root.get("genotype", ""),
        ]
        seg_map = {s.get("label"): s for s in root.get("segments", [])}
        for col in segment_cols:
            s = seg_map.get(col)
            row.append(s.get("length_px") if s else "")
            row.append(s.get("length_physical") if s else "")
        row += [
            root.get("total_length_px", ""),
            root.get("total_length_physical", ""),
            root.get("physical_unit", ""),
            root.get("tortuosity", ""),
            root.get("angle_deg", ""),
            root.get("deviation_deg", ""),
        ]
        ws.append(row)

    for i, _ in enumerate(headers, 1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "AA"].width = 16

    # --- Polyline sheet (one row per vertex, for reproducibility) ---
    ws_pl = wb.create_sheet("Polylines")
    ws_pl.append(["Root ID", "Genotype", "Vertex", "X", "Y"])
    for cell in ws_pl[1]:
        cell.font = bold
    for root in data.get("roots", []):
        rid = root.get("root_id", "")
        geno = root.get("genotype", "")
        for i, pt in enumerate(root.get("polyline", [])):
            ws_pl.append([rid, geno, i, pt.get("x"), pt.get("y")])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"root_growth_{image_id}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
